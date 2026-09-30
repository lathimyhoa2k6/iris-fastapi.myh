"""Read the PostgreSQL database of the deployed service from its External Database URL.

Read-only: the session runs with default_transaction_read_only=on, so nothing can
be changed from here.

Usage (URL from --url, EXTERNAL_DATABASE_URL in the environment or .env, or typed
in a hidden prompt):
    python db_viewer.py                      # tables + row counts + latest predictions
    python db_viewer.py history --watch 5    # prediction history, refreshed every 5 s
    python db_viewer.py show users           # one table
    python db_viewer.py query "SELECT ..."   # any read-only SQL
    python db_viewer.py export               # every table to an .xlsx file
"""

from __future__ import annotations

import argparse
import getpass
import json
import os
import sys
import time
import unicodedata
from datetime import datetime, timedelta, timezone
from pathlib import Path

import psycopg
from psycopg import sql

BASE_DIR = Path(__file__).resolve().parent
UTC_FORMAT = "%Y-%m-%dT%H:%M:%SZ"
TZ = timezone(timedelta(hours=7))
MAX_CELL = 90

HISTORY_SQL = """
    SELECT p.id, u.username, p.created_at, p.task, p.model, p.input_json,
           p.predicted_label, p.predicted_value, p.actual_value, p.runtime_ms
    FROM predictions p JOIN users u ON u.id = p.user_id
    {where}
    ORDER BY p.id DESC LIMIT %s
"""


# ------------------------------------------------------------------ connection


def _dotenv_value(name: str) -> str | None:
    path = BASE_DIR / ".env"
    if not path.exists():
        return None
    for line in path.read_text(encoding="utf-8").splitlines():
        key, sep, value = line.strip().partition("=")
        if sep and key.strip() == name:
            return value.strip().strip('"').strip("'") or None
    return None


def resolve_url(cli_url: str | None) -> str:
    url = cli_url or os.environ.get("EXTERNAL_DATABASE_URL") or _dotenv_value("EXTERNAL_DATABASE_URL")
    if not url:
        print("Dán External Database URL (Render → iris-svm-db → Connect → External).")
        url = getpass.getpass("URL (ẩn khi gõ): ").strip()
    if not url.startswith(("postgres://", "postgresql://")):
        sys.exit("URL không hợp lệ: phải bắt đầu bằng postgresql:// hoặc postgres://")
    return url


def connect(url: str) -> psycopg.Connection:
    try:
        return psycopg.connect(url, autocommit=True, connect_timeout=20,
                               options="-c default_transaction_read_only=on")
    except psycopg.OperationalError as exc:
        # The message never contains the password.
        sys.exit(f"Không kết nối được CSDL: {str(exc).strip().splitlines()[0]}")


def host_label(conn: psycopg.Connection) -> str:
    info = conn.info
    return f"{info.dbname}@{info.host}"


# ------------------------------------------------------------------ formatting


def to_local(value):
    """ISO UTC text timestamps ('...Z') -> local time (UTC+7) for display."""
    if isinstance(value, str) and len(value) == 20 and value.endswith("Z"):
        try:
            return datetime.strptime(value, UTC_FORMAT).replace(tzinfo=timezone.utc) \
                .astimezone(TZ).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            return value
    return value


def _width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        value = f"{value:.4g}"
    text = str(to_local(value)).replace("\n", " ")
    return text if len(text) <= MAX_CELL else text[: MAX_CELL - 1] + "…"


def _short_input(text):
    """'{"sepal_length": 5.1, ...}' -> 'sepal_length=5.1, ...' so the four inputs fit on screen."""
    try:
        return ", ".join(f"{k}={v}" for k, v in json.loads(text).items())
    except (TypeError, ValueError, AttributeError):
        return text


def print_table(columns: list[str], rows: list[tuple]) -> None:
    short = {i for i, c in enumerate(columns) if c == "input_json"}
    cells = [[_cell(_short_input(v) if i in short else v) for i, v in enumerate(row)] for row in rows]
    widths = [max([_width(c)] + [_width(r[i]) for r in cells]) for i, c in enumerate(columns)]

    def line(values):
        return " | ".join(v + " " * (w - _width(v)) for v, w in zip(values, widths))

    print(line(columns))
    print("-+-".join("-" * w for w in widths))
    for row in cells:
        print(line(row))
    print(f"({len(rows)} dòng)")


def run_and_print(conn, query, params=None) -> None:
    # params=None keeps a literal % in hand-written SQL (e.g. LIKE 'a%') as is.
    cur = conn.execute(query, params)
    if cur.description is None:
        print("(câu lệnh không trả về dữ liệu)")
        return
    print_table([c.name for c in cur.description], cur.fetchall())


# ------------------------------------------------------------------ commands


def table_names(conn) -> list[str]:
    return [r[0] for r in conn.execute(
        "SELECT table_name FROM information_schema.tables "
        "WHERE table_schema = 'public' AND table_type = 'BASE TABLE' ORDER BY table_name")]


def cmd_overview(conn, args) -> None:
    print(f"CSDL: {host_label(conn)}\n")
    counts = [(t, conn.execute(sql.SQL("SELECT COUNT(*) FROM {}").format(sql.Identifier(t))).fetchone()[0])
              for t in table_names(conn)]
    print_table(["bảng", "số dòng"], counts)
    print("\n10 dự đoán mới nhất:")
    run_and_print(conn, HISTORY_SQL.format(where=""), (10,))


def cmd_history(conn, args) -> None:
    where, params = "", []
    if args.user:
        where = "WHERE lower(u.username) = lower(%s)"
        params.append(args.user)
    query = HISTORY_SQL.format(where=where)
    while True:
        if args.watch:
            os.system("cls" if os.name == "nt" else "clear")
            print(f"CSDL: {host_label(conn)} · tự làm mới mỗi {args.watch}s · "
                  f"{datetime.now(TZ):%H:%M:%S} · Ctrl+C để thoát\n")
        run_and_print(conn, query, (*params, args.limit))
        if not args.watch:
            return
        sys.stdout.flush()
        time.sleep(args.watch)


def cmd_show(conn, args) -> None:
    names = table_names(conn)
    if args.table not in names:
        sys.exit(f"Không có bảng '{args.table}'. Các bảng: {', '.join(names)}")
    run_and_print(conn, sql.SQL("SELECT * FROM {} ORDER BY 1 DESC LIMIT %s").format(sql.Identifier(args.table)),
                  (args.limit,))


def cmd_query(conn, args) -> None:
    try:
        # prepare=True: exactly one statement, so "SET ...; DELETE ..." cannot slip through.
        cur = conn.execute(args.sql, prepare=True)
        if cur.description is None:
            print("(câu lệnh không trả về dữ liệu)")
            return
        print_table([c.name for c in cur.description], cur.fetchall())
    except psycopg.errors.ReadOnlySqlTransaction:
        sys.exit("Công cụ này chỉ đọc: không chạy được INSERT/UPDATE/DELETE/DDL.")
    except psycopg.Error as exc:
        sys.exit(f"Lỗi SQL: {str(exc).strip()}")


def cmd_export(conn, args) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for name in table_names(conn):
        cur = conn.execute(sql.SQL("SELECT * FROM {} ORDER BY 1").format(sql.Identifier(name)))
        ws = wb.create_sheet(name[:31])
        ws.append([c.name for c in cur.description])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="6C5CE7")
        for row in cur:
            ws.append([to_local(v) for v in row])
        ws.freeze_panes = "A2"
        print(f"  {name}: {ws.max_row - 1} dòng")
    out = Path(args.out or f"csdl_{datetime.now(TZ):%Y%m%d_%H%M%S}.xlsx")
    wb.save(out)
    print(f"Đã xuất {out.resolve()}")


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Xem CSDL PostgreSQL trên Render bằng External Database URL (chỉ đọc).")
    parser.add_argument("--url", help="External Database URL (mặc định: EXTERNAL_DATABASE_URL hoặc hỏi ẩn)")
    sub = parser.add_subparsers(dest="command")

    history = sub.add_parser("history", help="lịch sử dự đoán, mới nhất trước")
    history.add_argument("--user", help="chỉ của một tài khoản")
    history.add_argument("--limit", type=int, default=20)
    history.add_argument("--watch", type=int, metavar="GIÂY", help="tự làm mới sau mỗi GIÂY giây")

    show = sub.add_parser("show", help="xem một bảng")
    show.add_argument("table", help="users | predictions | training_runs | model_runs | schema_migrations")
    show.add_argument("--limit", type=int, default=50)

    query = sub.add_parser("query", help="chạy một câu SQL chỉ đọc")
    query.add_argument("sql")

    export = sub.add_parser("export", help="xuất mọi bảng ra Excel")
    export.add_argument("--out", help="tên tệp .xlsx")

    args = parser.parse_args()
    commands = {"history": cmd_history, "show": cmd_show, "query": cmd_query, "export": cmd_export}
    with connect(resolve_url(args.url)) as conn:
        try:
            commands.get(args.command, cmd_overview)(conn, args)
        except KeyboardInterrupt:
            print("\nĐã dừng.")


if __name__ == "__main__":
    main()
